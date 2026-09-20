from __future__ import annotations

from collections.abc import Generator
from io import BytesIO
from pathlib import Path
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database.base import Base
from app.database.session import get_db_session
from app.main import app
from app.modules.profile.models import Education, Experience, Profile, Project, Skill
from app.modules.resume.models import Resume


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Generator[TestClient, None, None]:
    monkeypatch.setenv("AUTH_SECRET", "test-only-auth-secret-with-32-bytes")
    monkeypatch.setenv("RESUME_STORAGE_PATH", str(tmp_path / "resumes"))
    from app.core.config import get_settings

    get_settings.cache_clear()
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_session() -> Generator[Session, None, None]:
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db_session] = override_session
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()
    get_settings.cache_clear()


def register(client: TestClient, email: str = "user@example.com", password: str = "password123"):
    return client.post("/api/v1/auth/register", json={"email": email, "password": password})


def login(client: TestClient, email: str = "user@example.com", password: str = "password123"):
    response = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": password},
    )
    return response


def get_auth_header(client: TestClient, email: str = "user@example.com") -> dict[str, str]:
    login_response = login(client, email=email)
    assert login_response.status_code == 200
    return {"Authorization": f"Bearer {login_response.json()['access_token']}"}


def build_pdf_bytes() -> bytes:
    return b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF"


def build_docx_bytes() -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\"></Types>")
        archive.writestr("word/document.xml", "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"></w:document>")
    return buffer.getvalue()


def test_authenticated_user_can_upload_valid_pdf(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    response = client.post(
        "/api/v1/resumes",
        files={"file": ("resume.pdf", build_pdf_bytes(), "application/pdf")},
        headers=headers,
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["original_filename"] == "resume.pdf"
    assert payload["mime_type"] == "application/pdf"
    assert payload["stored_filename"] != "resume.pdf"
    assert payload["storage_key"]


def test_authenticated_user_can_upload_valid_docx(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    response = client.post(
        "/api/v1/resumes",
        files={
            "file": (
                "resume.docx",
                build_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["original_filename"] == "resume.docx"
    assert (
        payload["mime_type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert payload["stored_filename"].endswith(".docx")


def test_unsupported_resume_type_is_rejected(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    response = client.post(
        "/api/v1/resumes",
        files={"file": ("resume.txt", b"plain text", "text/plain")},
        headers=headers,
    )

    assert response.status_code == 415


def test_oversized_resume_is_rejected(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    oversized = b"%PDF-1.4\n" + b"A" * (11 * 1024 * 1024)
    response = client.post(
        "/api/v1/resumes",
        files={"file": ("large.pdf", oversized, "application/pdf")},
        headers=headers,
    )

    assert response.status_code == 413


def test_unsafe_filename_path_is_handled_safely(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    response = client.post(
        "/api/v1/resumes",
        files={"file": ("../../evil.pdf", build_pdf_bytes(), "application/pdf")},
        headers=headers,
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["original_filename"] == "evil.pdf"
    assert ".." not in payload["storage_key"]
    assert payload["storage_key"].endswith(".pdf")


def test_user_can_list_only_their_own_resumes(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")

    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")

    owner_upload = client.post(
        "/api/v1/resumes",
        files={"file": ("owner.pdf", build_pdf_bytes(), "application/pdf")},
        headers=owner_headers,
    )
    assert owner_upload.status_code == 201

    other_list = client.get("/api/v1/resumes", headers=other_headers)
    assert other_list.status_code == 200
    assert other_list.json() == []

    owner_list = client.get("/api/v1/resumes", headers=owner_headers)
    assert owner_list.status_code == 200
    assert len(owner_list.json()) == 1


def test_user_can_access_their_own_resume_and_not_another_users(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")

    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("owner.pdf", build_pdf_bytes(), "application/pdf")},
        headers=owner_headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    own = client.get(f"/api/v1/resumes/{resume_id}", headers=owner_headers)
    assert own.status_code == 200
    assert own.json()["id"] == resume_id

    other = client.get(f"/api/v1/resumes/{resume_id}", headers=other_headers)
    assert other.status_code == 404


def test_user_can_delete_their_own_resume_and_storage_file(
    client: TestClient,
    tmp_path: Path,
) -> None:
    register(client)
    headers = get_auth_header(client)

    response = client.post(
        "/api/v1/resumes",
        files={"file": ("delete-me.pdf", build_pdf_bytes(), "application/pdf")},
        headers=headers,
    )
    assert response.status_code == 201
    payload = response.json()
    resume_id = payload["id"]
    storage_key = payload["storage_key"]
    storage_file = tmp_path / "resumes" / storage_key
    assert storage_file.exists()

    delete_response = client.delete(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert delete_response.status_code == 204
    assert not storage_file.exists()

    session = next(app.dependency_overrides[get_db_session]())
    assert session.scalar(select(Resume).where(Resume.id == UUID(resume_id))) is None
    session.close()


def test_user_cannot_delete_another_users_resume(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")

    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("owner.pdf", build_pdf_bytes(), "application/pdf")},
        headers=owner_headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    response = client.delete(f"/api/v1/resumes/{resume_id}", headers=other_headers)
    assert response.status_code == 404


def test_resume_routes_require_authentication(client: TestClient) -> None:
    assert (
        client.post(
            "/api/v1/resumes",
            files={"file": ("resume.pdf", build_pdf_bytes(), "application/pdf")},
        ).status_code
        == 401
    )
    assert client.get("/api/v1/resumes").status_code == 401
    assert client.get("/api/v1/resumes/123e4567-e89b-12d3-a456-426614174000").status_code == 401
    assert client.delete("/api/v1/resumes/123e4567-e89b-12d3-a456-426614174000").status_code == 401


def test_resume_is_linked_to_profile_owner(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    client.post(
        "/api/v1/resumes",
        files={"file": ("profile-bound.pdf", build_pdf_bytes(), "application/pdf")},
        headers=headers,
    )

    session = next(app.dependency_overrides[get_db_session]())
    profile = session.scalar(select(Profile))
    resume = session.scalar(select(Resume))
    assert profile is not None
    assert resume is not None
    assert resume.profile_id == profile.id
    session.close()


def test_resume_parser_extracts_basic_information_and_skills() -> None:
    from app.modules.resume.parser import ResumeParser

    text = """
    Jane Doe
    jane.doe@example.com | +1 (555) 123-4567
    Boston, MA
    linkedin.com/in/janedoe | github.com/janedoe

    Skills
    Python, SQL, Docker, FastAPI

    Programming Languages
    Python | Java | SQL
    """

    result = ResumeParser().parse(text)
    assert result.basic_information.full_name == "Jane Doe"
    assert result.basic_information.email == "jane.doe@example.com"
    assert result.basic_information.phone == "+1 (555) 123-4567"
    assert result.basic_information.location == "Boston, MA"
    assert result.basic_information.linkedin_url == "https://linkedin.com/in/janedoe"
    assert result.basic_information.github_url == "https://github.com/janedoe"
    assert "Python" in [skill.name for skill in result.skills]
    assert any(skill.category == "Programming Languages" for skill in result.skills)


def test_resume_parser_extracts_education_experience_and_projects() -> None:
    from app.modules.resume.parser import ResumeParser

    text = """
    Summary
    Experienced engineer.

    Education
    B.Tech in Computer Science, University of Technology, 2014 - 2018

    Experience
    Senior Software Engineer | Acme Corp | Boston, MA | 2020 - Present
    Built APIs and data pipelines.

    Projects
    JobPilot AI | Python, FastAPI, PostgreSQL | https://example.com/project
    Built a job-matching dashboard.
    """

    result = ResumeParser().parse(text)
    assert result.education[0].institution == "University of Technology"
    assert result.education[0].degree == "B.Tech"
    assert result.education[0].field_of_study == "Computer Science"
    assert result.experience[0].company == "Acme Corp"
    assert result.experience[0].job_title == "Senior Software Engineer"
    assert result.experience[0].end_date is None
    assert result.projects[0].name == "JobPilot AI"
    assert "Python" in result.projects[0].technologies


def test_resume_parser_handles_missing_sections_gracefully() -> None:
    from app.modules.resume.parser import ResumeParser

    result = ResumeParser().parse("No useful structured data here.")
    assert result.basic_information is None
    assert result.skills == []
    assert result.education == []
    assert result.experience == []
    assert result.projects == []
    assert result.warnings


def test_resume_parse_endpoint_requires_ownership(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")

    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")

    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Alice Developer\nAlice@example.com\nSkills\nPython\n")
    pdf_bytes = pdf.tobytes()
    pdf.close()

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("parser.pdf", pdf_bytes, "application/pdf")},
        headers=owner_headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    parsed = client.post(f"/api/v1/resumes/{resume_id}/parse", headers=other_headers)
    assert parsed.status_code == 404

    unauthenticated = client.post(f"/api/v1/resumes/{resume_id}/parse")
    assert unauthenticated.status_code == 401


def test_resume_parse_endpoint_extracts_structured_data(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "Jane Doe\nJane.Doe@example.com\n(555) 123-4567\nBoston, MA\n"
        "https://linkedin.com/in/janedoe\nhttps://github.com/janedoe\n\n"
        "Skills\nPython, SQL, FastAPI\n\nEducation\n"
        "B.Tech in Computer Science, University of Technology, 2014 - 2018\n\n"
        "Experience\nSenior Software Engineer | Acme Corp | Boston, MA | 2020 - Present\n"
        "Built APIs.\n",
    )
    pdf_bytes = pdf.tobytes()
    pdf.close()

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("structured.pdf", pdf_bytes, "application/pdf")},
        headers=headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    response = client.post(f"/api/v1/resumes/{resume_id}/parse", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["basic_information"]["full_name"] == "Jane Doe"
    assert payload["basic_information"]["email"] == "Jane.Doe@example.com"
    assert "Python" in [skill["name"] for skill in payload["skills"]]
    assert payload["education"][0]["institution"] == "University of Technology"
    assert payload["experience"][0]["company"] == "Acme Corp"


def build_review_pdf_bytes() -> bytes:
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        (72, 72),
        "Jane Doe\njane@example.com\nBoston, MA\n\nSkills\nPython, SQL\n\n"
        "Education\nB.Tech in Computer Science, University, 2014 - 2018\n\n"
        "Experience\nEngineer | Acme | Boston, MA | 2020 - 2022\nBuilt systems.\n\n"
        "Projects\nResume Tool | Python | https://example.com",
    )
    value = pdf.tobytes()
    pdf.close()
    return value


def create_review_resume(client: TestClient, headers: dict[str, str]) -> str:
    created = client.post(
        "/api/v1/resumes",
        files={"file": ("review.pdf", build_review_pdf_bytes(), "application/pdf")},
        headers=headers,
    )
    assert created.status_code == 201
    return created.json()["id"]


def test_resume_review_requires_authentication_and_ownership(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")
    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")
    resume_id = create_review_resume(client, owner_headers)

    assert client.get(f"/api/v1/resumes/{resume_id}/review").status_code == 401
    assert client.post(f"/api/v1/resumes/{resume_id}/approve").status_code == 401
    assert (
        client.get(f"/api/v1/resumes/{resume_id}/review", headers=other_headers).status_code
        == 404
    )
    assert (
        client.post(f"/api/v1/resumes/{resume_id}/approve", headers=other_headers).status_code
        == 404
    )

    response = client.get(f"/api/v1/resumes/{resume_id}/review", headers=owner_headers)
    assert response.status_code == 200
    assert response.json()["status"] == "PARSED"
    assert response.json()["basic_information"]["full_name"] == "Jane Doe"


def test_resume_review_rejects_invalid_edits_and_accepts_owner_edits(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)
    resume_id = create_review_resume(client, headers)

    invalid = client.put(
        f"/api/v1/resumes/{resume_id}/review",
        json={"skills": [{"name": ""}]},
        headers=headers,
    )
    assert invalid.status_code == 422

    edited = client.put(
        f"/api/v1/resumes/{resume_id}/review",
        json={
            "basic_information": {"full_name": "Edited Name", "email": "edited@example.com"},
            "skills": [{"name": "Python", "category": "Languages"}],
            "education": [],
            "experience": [],
            "projects": [],
            "warnings": [],
        },
        headers=headers,
    )
    assert edited.status_code == 200
    assert edited.json()["basic_information"]["full_name"] == "Edited Name"


def test_unapproved_review_does_not_write_trusted_data(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)
    resume_id = create_review_resume(client, headers)
    response = client.get(f"/api/v1/resumes/{resume_id}/review", headers=headers)
    assert response.status_code == 200

    session = next(app.dependency_overrides[get_db_session]())
    assert session.scalars(select(Skill)).all() == []
    assert session.scalars(select(Education)).all() == []
    assert session.scalars(select(Experience)).all() == []
    assert session.scalars(select(Project)).all() == []
    session.close()


def test_approval_uses_edits_and_is_idempotent(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)
    resume_id = create_review_resume(client, headers)
    edited = client.put(
        f"/api/v1/resumes/{resume_id}/review",
        json={
            "basic_information": {"full_name": "Approved Name", "email": "approved@example.com"},
            "skills": [{"name": "Python", "category": "Languages"}],
            "education": [
                {
                    "institution": "University",
                    "degree": "B.Tech",
                    "field_of_study": "Computer Science",
                    "start_date": "2014-01-01",
                    "end_date": "2018-01-01",
                }
            ],
            "experience": [
                {
                    "company": "Edited Co",
                    "job_title": "Engineer",
                    "start_date": "2020-01-01",
                    "end_date": "2022-01-01",
                }
            ],
            "projects": [
                {
                    "name": "Edited Project",
                    "start_date": "2021-01-01",
                    "technologies": ["Python"],
                }
            ],
            "warnings": [],
        },
        headers=headers,
    )
    assert edited.status_code == 200

    approved = client.post(f"/api/v1/resumes/{resume_id}/approve", headers=headers)
    assert approved.status_code == 200
    assert approved.json()["status"] == "APPROVED"

    session = next(app.dependency_overrides[get_db_session]())
    profile = session.scalar(select(Profile))
    assert profile is not None
    assert profile.full_name == "Approved Name"
    assert session.scalars(select(Skill)).all()[0].name == "Python"
    assert session.scalars(select(Experience)).all()[0].company == "Edited Co"
    assert session.scalars(select(Project)).all()[0].name == "Edited Project"
    assert len(session.scalars(select(Education)).all()) == 1
    session.close()

    repeated = client.post(f"/api/v1/resumes/{resume_id}/approve", headers=headers)
    assert repeated.status_code == 200
    session = next(app.dependency_overrides[get_db_session]())
    assert len(session.scalars(select(Skill)).all()) == 1
    assert len(session.scalars(select(Education)).all()) == 1
    assert len(session.scalars(select(Experience)).all()) == 1
    assert len(session.scalars(select(Project)).all()) == 1
    session.close()
    register(client)
    headers = get_auth_header(client)

    import fitz

    pdf_bytes = fitz.open()
    page = pdf_bytes.new_page()
    page.insert_text((72, 72), "Candidate Name\nSenior Python Engineer")
    pdf_buffer = pdf_bytes.tobytes()
    pdf_bytes.close()

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("candidate.pdf", pdf_buffer, "application/pdf")},
        headers=headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    response = client.post(f"/api/v1/resumes/{resume_id}/process", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["resume_id"] == resume_id
    assert payload["status"] == "success"
    assert "Candidate Name" in payload["text"]
    assert payload["character_count"] > 0
    assert payload["extractor"] == "pdf"


def test_resume_processing_extracts_docx_paragraphs(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    from docx import Document

    document = Document()
    document.add_paragraph("Alice Candidate")
    document.add_paragraph("Full Stack Engineer")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    doc_buffer = BytesIO()
    document.save(doc_buffer)

    created = client.post(
        "/api/v1/resumes",
        files={
            "file": (
                "candidate.docx",
                doc_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    response = client.post(f"/api/v1/resumes/{resume_id}/process", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert "Alice Candidate" in payload["text"]
    assert "Python" in payload["text"]
    assert payload["extractor"] == "docx"


def test_resume_processing_handles_image_only_pdf(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)

    pdf_bytes = build_pdf_bytes()
    created = client.post(
        "/api/v1/resumes",
        files={"file": ("scanned.pdf", pdf_bytes, "application/pdf")},
        headers=headers,
    )
    assert created.status_code == 201
    resume_id = created.json()["id"]

    response = client.post(f"/api/v1/resumes/{resume_id}/process", headers=headers)
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "error"
    assert payload["error_code"] == "OCR_REQUIRED"
    assert "OCR" in payload["error_message"]


def test_resume_processing_requires_ownership(client: TestClient) -> None:
    register(client, "owner@example.com")
    register(client, "other@example.com")

    owner_headers = get_auth_header(client, "owner@example.com")
    other_headers = get_auth_header(client, "other@example.com")

    created = client.post(
        "/api/v1/resumes",
        files={"file": ("owner.pdf", build_pdf_bytes(), "application/pdf")},
        headers=owner_headers,
    )
    resume_id = created.json()["id"]

    response = client.post(f"/api/v1/resumes/{resume_id}/process", headers=other_headers)
    assert response.status_code == 404


def test_resume_processing_requires_authentication(client: TestClient) -> None:
    register(client)
    headers = get_auth_header(client)
    created = client.post(
        "/api/v1/resumes",
        files={"file": ("resume.pdf", build_pdf_bytes(), "application/pdf")},
        headers=headers,
    )
    resume_id = created.json()["id"]

    assert client.post(f"/api/v1/resumes/{resume_id}/process").status_code == 401
