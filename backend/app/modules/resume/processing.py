"""Resume document extraction and normalization utilities."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Protocol
from zipfile import BadZipFile

import fitz  # type: ignore[import-untyped]
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


class ExtractionError(RuntimeError):
    """Raised when a supported document cannot be extracted safely."""

    def __init__(self, message: str, code: str) -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass(slots=True)
class ProcessingResult:
    """Internal processing outcome for a stored resume document."""

    success: bool
    text: str = ""
    character_count: int = 0
    page_count: int | None = None
    extractor: str = ""
    error_code: str | None = None
    error_message: str | None = None

    @property
    def status(self) -> str:
        return "success" if self.success else "error"


class DocumentExtractor(Protocol):
    """Common interface for format-specific text extractors."""

    name: str

    def extract(self, file_bytes: bytes) -> str:
        """Return normalized plain-text content extracted from the file."""


class PDFExtractor:
    """Extract textual content from a PDF file when the library can read actual text."""

    name = "pdf"

    def extract(self, file_bytes: bytes) -> str:
        try:
            document = fitz.open(stream=file_bytes, filetype="pdf")
        except (RuntimeError, ValueError, TypeError) as exc:
            raise ExtractionError("Corrupted PDF file.", "CORRUPTED_PDF") from exc

        try:
            pages: list[str] = []
            for page in document:
                page_text = page.get_text("text").strip()
                if page_text:
                    pages.append(page_text)
            full_text = "\n\n".join(pages)
            if not full_text.strip():
                raise ExtractionError("No extractable text found. OCR is required.", "OCR_REQUIRED")
            return full_text
        except Exception as exc:
            if isinstance(exc, ExtractionError):
                raise
            raise ExtractionError("Failed to extract text from PDF.", "EXTRACTION_FAILED") from exc
        finally:
            document.close()


class DOCXExtractor:
    """Extract paragraphs and table text from a DOCX document."""

    name = "docx"

    def extract(self, file_bytes: bytes) -> str:
        try:
            document = Document(BytesIO(file_bytes))
        except (BadZipFile, PackageNotFoundError, ValueError, TypeError) as exc:
            raise ExtractionError("Corrupted DOCX file.", "CORRUPTED_DOCX") from exc

        try:
            text_blocks: list[str] = []
            for paragraph in document.paragraphs:
                value = paragraph.text.strip()
                if value:
                    text_blocks.append(value)

            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    cleaned = [value for value in values if value]
                    if cleaned:
                        text_blocks.append(" | ".join(cleaned))

            text = "\n".join(text_blocks)
            if not text.strip():
                raise ExtractionError("Document contains no extractable text.", "EMPTY_DOCUMENT")
            return text
        except Exception as exc:
            if isinstance(exc, ExtractionError):
                raise
            raise ExtractionError("Failed to extract text from DOCX.", "EXTRACTION_FAILED") from exc


def normalize_text(raw_text: str) -> str:
    """Normalize line endings and whitespace while preserving paragraph boundaries."""

    normalized = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    lines: list[str] = []
    previous_blank = False

    for line in normalized.split("\n"):
        cleaned = re.sub(r"\s+", " ", line).strip()
        if not cleaned:
            if lines and not previous_blank:
                lines.append("")
            previous_blank = True
            continue
        lines.append(cleaned)
        previous_blank = False

    if lines and lines[-1] == "":
        lines = lines[:-1]
    text = "\n".join(lines).strip()
    return text.encode("utf-8", errors="replace").decode("utf-8")


def get_document_extractor(mime_type: str, filename: str) -> DocumentExtractor:
    """Select an extractor based on the known resume format."""

    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if mime_type == "application/pdf" or suffix == "pdf":
        return PDFExtractor()
    if (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == "docx"
    ):
        return DOCXExtractor()
    raise ExtractionError("Unsupported resume format.", "UNSUPPORTED_FORMAT")
